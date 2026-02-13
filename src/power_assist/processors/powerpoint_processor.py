"""PowerPoint document processor"""
import os
from typing import Dict, Any
from ..models.document import DocumentContent, DocumentType
from .base import BaseProcessor


class PowerPointProcessor(BaseProcessor):
    """Processor for PowerPoint documents"""
    
    @property
    def supported_type(self) -> DocumentType:
        return DocumentType.POWERPOINT
    
    async def validate(self, file_path: str) -> bool:
        """Validate PowerPoint file"""
        if not os.path.exists(file_path):
            return False
        return file_path.lower().endswith(('.ppt', '.pptx'))
    
    async def process(self, file_path: str) -> DocumentContent:
        """
        Process PowerPoint document
        
        Args:
            file_path: Path to PowerPoint file
            
        Returns:
            DocumentContent with extracted information
        """
        try:
            from pptx import Presentation
            
            prs = Presentation(file_path)
            text_parts = []
            
            for slide_num, slide in enumerate(prs.slides, 1):
                text_parts.append(f"## Slide {slide_num}")
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text_parts.append(shape.text)
                
                text_parts.append("")  # Empty line between slides
            
            metadata: Dict[str, Any] = {
                "num_slides": len(prs.slides),
                "source": file_path,
                "processor": "python-pptx",
            }
            
            return DocumentContent(
                text="\n".join(text_parts),
                metadata=metadata,
            )
            
        except Exception as e:
            raise Exception(f"PowerPoint processing failed: {str(e)}")
