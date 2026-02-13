"""Word document processor"""
import os
from typing import Dict, Any
from ..models.document import DocumentContent, DocumentType
from .base import BaseProcessor


class WordProcessor(BaseProcessor):
    """Processor for Word documents"""
    
    @property
    def supported_type(self) -> DocumentType:
        return DocumentType.WORD
    
    async def validate(self, file_path: str) -> bool:
        """Validate Word file"""
        if not os.path.exists(file_path):
            return False
        return file_path.lower().endswith(('.doc', '.docx'))
    
    async def process(self, file_path: str) -> DocumentContent:
        """
        Process Word document
        
        Args:
            file_path: Path to Word file
            
        Returns:
            DocumentContent with extracted information
        """
        try:
            from docx import Document
            
            doc = Document(file_path)
            text_parts = []
            
            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Extract tables
            tables = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables.append({"data": table_data})
            
            metadata: Dict[str, Any] = {
                "num_paragraphs": len(doc.paragraphs),
                "num_tables": len(doc.tables),
                "source": file_path,
                "processor": "python-docx",
            }
            
            return DocumentContent(
                text="\n\n".join(text_parts),
                metadata=metadata,
                tables=tables,
            )
            
        except Exception as e:
            raise Exception(f"Word processing failed: {str(e)}")
